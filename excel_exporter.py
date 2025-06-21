import pandas as pd
import calendar
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

class ExcelExporter:
    TURKISH_MONTHS = {
        1: "Ocak", 2: "Şubat", 3: "Mart", 4: "Nisan", 5: "Mayıs", 6: "Haziran",
        7: "Temmuz", 8: "Ağustos", 9: "Eylül", 10: "Ekim", 11: "Kasım", 12: "Aralık"
    }
    
    def __init__(self, database):
        self.db = database
    
    def export_monthly_schedule(self, year, month, filename=None):
        if not filename:
            filename = f"nobet_programi_{year}_{month:02d}.xlsx"
        
        schedule_data = self.get_monthly_schedule(year, month)
        
        wb = Workbook()
        ws = wb.active
        ws.title = f"{year} {self.TURKISH_MONTHS[month]}"
        
        days_in_month = calendar.monthrange(year, month)[1]
        first_weekday = calendar.monthrange(year, month)[0]
        
        weekdays = ['Pazartesi', 'Salı', 'Çarşamba', 'Perşembe', 'Cuma', 'Cumartesi', 'Pazar']
        
        for col, day_name in enumerate(weekdays, 1):
            ws.cell(row=1, column=col, value=day_name)
            ws.cell(row=1, column=col).font = Font(bold=True)
            ws.cell(row=1, column=col).alignment = Alignment(horizontal='center')
        
        current_row = 2
        current_col = first_weekday + 1
        
        for day in range(1, days_in_month + 1):
            current_date = date(year, month, day)
            
            cell_value = f"{day}"
            if current_date in schedule_data:
                person_info = schedule_data[current_date]
                cell_value += f"\n{person_info['name']}\n({person_info['status']})"
            
            cell = ws.cell(row=current_row, column=current_col, value=cell_value)
            cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
            
            if current_date in schedule_data:
                cell.fill = PatternFill(start_color="E6F3FF", end_color="E6F3FF", fill_type="solid")
            
            current_col += 1
            if current_col > 7:
                current_col = 1
                current_row += 1
        
        for col in range(1, 8):
            ws.column_dimensions[chr(64 + col)].width = 15
        
        for row in range(2, current_row + 1):
            ws.row_dimensions[row].height = 60
        
        wb.save(filename)
        return filename
    
    def get_monthly_schedule(self, year, month):
        conn = self.db.get_connection()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT n.tarih, p.ad, p.statu 
            FROM Nobet n
            JOIN Personel p ON n.personelId = p.id
            WHERE strftime('%Y-%m', n.tarih) = ?
            ORDER BY n.tarih
        """, (f"{year:04d}-{month:02d}",))
        
        results = cursor.fetchall()
        conn.close()
        
        schedule_dict = {}
        for row in results:
            schedule_date = date.fromisoformat(row[0])
            schedule_dict[schedule_date] = {
                'name': row[1],
                'status': row[2]
            }
        
        return schedule_dict
    
    def export_in_memory_schedule(self, schedule_data, year, month, filename=None):
        if not filename:
            filename = f"nobet_programi_{year}_{month:02d}.xlsx"
        
        schedule_dict = {}
        for scheduled_date, person_id, day_type in schedule_data:
            if person_id:
                conn = self.db.get_connection()
                cursor = conn.cursor()
                cursor.execute("SELECT ad, statu FROM Personel WHERE id = ?", (person_id,))
                person_info = cursor.fetchone()
                conn.close()
                
                if person_info:
                    schedule_dict[scheduled_date] = {
                        'name': person_info[0],
                        'status': person_info[1]
                    }
        
        wb = Workbook()
        ws = wb.active
        ws.title = f"{year} {self.TURKISH_MONTHS[month]}"
        
        days_in_month = calendar.monthrange(year, month)[1]
        first_weekday = calendar.monthrange(year, month)[0]
        
        weekdays = ['Pazartesi', 'Salı', 'Çarşamba', 'Perşembe', 'Cuma', 'Cumartesi', 'Pazar']
        
        for col, day_name in enumerate(weekdays, 1):
            ws.cell(row=1, column=col, value=day_name)
            ws.cell(row=1, column=col).font = Font(bold=True)
            ws.cell(row=1, column=col).alignment = Alignment(horizontal='center')
        
        current_row = 2
        current_col = first_weekday + 1
        
        for day in range(1, days_in_month + 1):
            current_date = date(year, month, day)
            
            cell_value = f"{day}"
            if current_date in schedule_dict:
                person_info = schedule_dict[current_date]
                cell_value += f"\n{person_info['name']}\n({person_info['status']})"
            
            cell = ws.cell(row=current_row, column=current_col, value=cell_value)
            cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
            
            if current_date in schedule_dict:
                cell.fill = PatternFill(start_color="E6F3FF", end_color="E6F3FF", fill_type="solid")
            
            current_col += 1
            if current_col > 7:
                current_col = 1
                current_row += 1
        
        for col in range(1, 8):
            ws.column_dimensions[chr(64 + col)].width = 15
        
        for row in range(2, current_row + 1):
            ws.row_dimensions[row].height = 60
        
        wb.save(filename)
        return filename
    
    def export_personnel_summary(self, year, month, filename=None):
        if not filename:
            filename = f"personel_ozet_{year}_{month:02d}.xlsx"
        
        conn = self.db.get_connection()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT 
                p.ad as 'Personel Adı',
                p.statu as 'Statü',
                COUNT(n.id) as 'Nöbet Sayısı',
                COALESCE(SUM(n.deger), 0) as 'Toplam Puan'
            FROM Personel p
            LEFT JOIN Nobet n ON p.id = n.personelId 
                AND strftime('%Y-%m', n.tarih) = ?
            WHERE p.Aktif = 1
            GROUP BY p.id, p.ad, p.statu
            ORDER BY p.ad
        """, (f"{year:04d}-{month:02d}",))
        
        data = cursor.fetchall()
        conn.close()
        
        df = pd.DataFrame(data, columns=['Personel Adı', 'Statü', 'Nöbet Sayısı', 'Toplam Puan'])
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Personel Özeti"
        
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        for col in range(1, len(df.columns) + 1):
            ws.cell(row=1, column=col).font = Font(bold=True)
            ws.column_dimensions[chr(64 + col)].width = 15
        
        wb.save(filename)
        return filename
